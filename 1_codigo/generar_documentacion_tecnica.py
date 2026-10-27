"""Genera Documentacion_Tecnica_ViT_DenseNet.docx con la misma estructura y formato que
Documentacion_Tecnica_10_Modelos.docx (repo Mendeley_ViT_DenseNet): metodologia por etapas con
figura de flujo y arquitectura, glosario de formulas, detalle por modelo con resultados por fold,
curvas ROC y calculo pooled, y anexos de trazabilidad y glosario de terminos.

Todos los numeros salen de los CSV de 3_resultados/; las graficas por fold se generan aqui
en 4_graficas/{trucha,rohu}_sinx2/{modelo}/.

Uso (desde la raiz del repo; requiere python-docx, pandas y matplotlib):
    python 1_codigo/generar_documentacion_tecnica.py [--plantilla RUTA_DOCX]
--plantilla: docx del que se toman estilos, margenes y tamano de pagina (por defecto la
documentacion tecnica de Mendeley, en el repo hermano).
"""
import argparse
import os
import statistics

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from matplotlib.patches import FancyBboxPatch
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

OUT = "Documentacion_Tecnica_ViT_DenseNet.docx"
PLANTILLA = "../Mendeley_ViT_DenseNet/6_documentacion/Documentacion_Tecnica_10_Modelos.docx"
MODELOS = [  # (clave, nombre, familia, parametros, imagen, batch, cabeza entrenable)
    ("vit_base_patch16_224", "ViT-Base/16", "Transformer (ViT)", "85.8M", 224, 8, 768 * 2 + 2),
    ("densenet121", "DenseNet-121", "CNN", "7.0M", 256, 16, 1024 * 2 + 2),
]
# F1 macro pooled en Mendeley (dominio fuente): Mendeley_ViT_DenseNet/3_resultados_test_por_modelo/
# metricas_pooled_por_modelo.csv
F1_FUENTE = {"vit_base_patch16_224": 96.96, "densenet121": 95.66}
DOMINIOS = [("trucha", "Trucha"), ("rohu", "Rohu")]
LOG = {"trucha": "2_logs/trucha_congelado_sinx2.log", "rohu": "2_logs/rohu_congelado_sinx2.log"}

AZUL, GRIS, FORMULA = RGBColor(0x1F, 0x4E, 0x78), RGBColor(0x55, 0x55, 0x55), RGBColor(0x1F, 0x4E, 0x78)
CABECERA, RESALTE = "2A4D7A", "FFE699"
MAR_LR, MAR_TB = 70, 45


# --- datos -------------------------------------------------------------------
def cargar():
    datos = {}
    for dom, _ in DOMINIOS:
        resumen = pd.read_csv(f"3_resultados/{dom}/confmat_roc/resumen_confmat_roc.csv")
        oficial = pd.read_csv(f"3_resultados/{dom}/congelado_sinx2_resumen.csv")
        for clave, *_ in MODELOS:
            base = f"3_resultados/{dom}/confmat_roc/{clave}"
            filas = resumen[resumen.model == clave].set_index("fold")
            folds = []
            for k in range(1, 6):
                cm = pd.read_csv(f"{base}/confusion_matrix_fold{k}.csv", index_col=0).values
                fila = filas.loc[str(k)]
                folds.append({"cm": cm, "auc": float(fila.auc), "epochs": int(fila.epochs_run),
                              "roc": pd.read_csv(f"{base}/roc_curve_fold{k}.csv")})
                assert abs(metricas(cm)["f1"] - float(fila.f1)) < 0.01, (dom, clave, k)
            pooled = pd.read_csv(f"{base}/confusion_matrix_pooled.csv", index_col=0).values
            assert (pooled == sum(f["cm"] for f in folds)).all(), (dom, clave)  # pooled = suma de folds
            datos[dom, clave] = {
                "folds": folds, "pooled": pooled, "auc": float(filas.loc["pooled"].auc),
                "roc": pd.read_csv(f"{base}/roc_curve_pooled.csv"),
                "f1_oficial": float(oficial[oficial.model == clave].f1_mean.iloc[0]),
            }
    return datos


def metricas(cm):
    tp, fn, fp, tn = cm[0, 0], cm[0, 1], cm[1, 0], cm[1, 1]
    p_f, r_f = tp / (tp + fp) * 100, tp / (tp + fn) * 100
    p_n, r_n = tn / (tn + fn) * 100, tn / (tn + fp) * 100
    f1_f, f1_n = 2 * p_f * r_f / (p_f + r_f), 2 * p_n * r_n / (p_n + r_n)
    return {"tp": tp, "fn": fn, "fp": fp, "tn": tn, "n": cm.sum(), "p_f": p_f, "r_f": r_f, "p_n": p_n,
            "r_n": r_n, "f1_f": f1_f, "f1_n": f1_n, "f1": (f1_f + f1_n) / 2, "acc": (tp + tn) / cm.sum() * 100,
            "prec": (p_f + p_n) / 2, "rec": (r_f + r_n) / 2}


def indicadores(datos, dom, clave):
    f1_t, f1_s = datos[dom, clave]["f1_oficial"], F1_FUENTE[clave]
    gap = (f1_s - f1_t) / 100
    base = "densenet121"
    gap_base = (F1_FUENTE[base] - datos[dom, base]["f1_oficial"]) / 100
    return f1_t / f1_s * 100, gap, (gap_base - gap) / gap_base * 100


# --- graficas por fold (mismo estilo que Mendeley: cm_folds_* y roc_folds_*) --
def graficas_folds(datos, dom, clave, nombre, dominio):
    carpeta = f"4_graficas/{dom}_sinx2/{clave}"
    folds = datos[dom, clave]["folds"]
    fig, axes = plt.subplots(1, 5, figsize=(20, 4.6))
    for k, (ax, f) in enumerate(zip(axes, folds), start=1):
        cm = f["cm"]
        ax.imshow(cm, cmap="Blues")
        for i in range(2):
            for j in range(2):
                ax.text(j, i, cm[i, j], ha="center", va="center",
                        color="white" if cm[i, j] > cm.max() / 2 else "black")
        ax.set_xticks([0, 1], ["fresco", "no_fresco"])
        ax.set_yticks([0, 1], ["fresco", "no_fresco"])
        ax.set_xlabel("Prediccion")
        ax.set_title(f"Fold {k}\nF1={metricas(cm)['f1']:.2f}", fontsize=10)
    axes[0].set_ylabel("Real")
    fig.suptitle(f"{nombre} — {dominio} — Matrices de confusión por fold", fontsize=13)
    fig.tight_layout()
    fig.savefig(f"{carpeta}/confusion_matrix_folds.png", dpi=150)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6, 6))
    for k, f in enumerate(folds, start=1):
        ax.plot(f["roc"].fpr, f["roc"].tpr, lw=1.3, alpha=0.8, label=f"Fold {k} (AUC={f['auc']:.4f})")
    roc = datos[dom, clave]["roc"]
    ax.plot(roc.fpr, roc.tpr, color="black", lw=2.5, label=f"Pooled (AUC={datos[dom, clave]['auc']:.4f})")
    ax.plot([0, 1], [0, 1], "--", color="gray", lw=1)
    ax.set_xlabel("Tasa de falsos positivos")
    ax.set_ylabel("Tasa de verdaderos positivos")
    ax.set_title(f"{nombre} — {dominio} — Curvas ROC: 5 folds + pooled")
    ax.legend(loc="lower right", fontsize=8.5)
    fig.tight_layout()
    fig.savefig(f"{carpeta}/roc_curve_folds.png", dpi=150)
    plt.close(fig)
    return f"{carpeta}/confusion_matrix_folds.png", f"{carpeta}/roc_curve_folds.png"


ETAPAS = [  # (titulo, linea 1, linea 2) -- misma figura de flujo que la documentacion de Mendeley
    ("Dominios objetivo y preparación de datos", "Trucha: 1,150 imágenes (481 fresco / 669 no_fresco)",
     "Rohu: 1,261 imágenes (422 fresco / 839 no_fresco)"),
    ("Validación cruzada 5-fold en cada dominio", "Trucha: StratifiedKFold · Rohu: StratifiedGroupKFold",
     "por corrida: 4 folds entrenan · 1 fold evalúa (semilla 42)"),
    ("Configuración de los 2 modelos (congelado)", "ViT-Base/16 + DenseNet-121 desde el checkpoint de Mendeley",
     "backbone congelado · solo se entrena la cabeza"),
    ("Entrenamiento de la cabeza y mejor epoch", "Máx. 30 epochs · early stopping = 8",
     "criterio: mejor val_F1 macro por fold"),
    ("Evaluación por fold", "Mejor epoch → fold de evaluación",
     "matriz de confusión + probabilidades + métricas"),
    ("Consolidación OOF / pooled", "Concatenación de los 5 folds de evaluación",
     "1,150 (Trucha) y 1,261 (Rohu) predicciones out-of-fold"),
    ("Comparación e indicadores de robustez", "F1 macro · Accuracy · AUC · estabilidad",
     "TRD · Gap · RBE frente a Mendeley (dominio fuente)"),
]


VERIF = "3_resultados/{dom}/verificacion_10fold"
PACIENCIA, MAX_EPOCHS = 8, 30


def epoch_elegido(epochs_run):
    """Epoch del mejor val_F1 deducido del early stopping de finetune_eval_fold (corta con paciencia >= 8)."""
    return str(epochs_run - PACIENCIA) if epochs_run < MAX_EPOCHS else f"{MAX_EPOCHS - PACIENCIA}–{MAX_EPOCHS} (no registrado)"


def curvas_epoch(dom, clave, nombre, dominio):
    """val_F1 por epoch de la corrida de verificacion (10 folds); el punto marca el epoch elegido."""
    base = f"{VERIF.format(dom=dom)}/{clave}"
    folds = [pd.read_csv(f"{base}/epochs_fold{k}.csv") for k in range(1, 11)]
    fig, ax = plt.subplots(figsize=(9, 4.6))
    for k, e in enumerate(folds, start=1):
        mejor = e.loc[e.val_f1.idxmax()]
        linea, = ax.plot(e.epoch, e.val_f1, lw=1.2, label=f"Fold {k}")
        ax.plot(mejor.epoch, mejor.val_f1, "o", color=linea.get_color(), ms=6)
    ax.set_xlabel("Epoch")
    ax.set_ylabel("val_F1 macro (%)")
    ax.set_title(f"{nombre} — {dominio} — val_F1 por epoch (corrida de verificación, 10 folds)")
    ax.legend(ncol=5, fontsize=8, loc="lower right")
    ax.grid(alpha=0.3)
    fig.tight_layout()
    ruta = f"4_graficas/{dom}_sinx2/{clave}/val_f1_por_epoch_10fold.png"
    fig.savefig(ruta, dpi=150)
    plt.close(fig)
    return ruta, folds


def figura_flujo(ruta="4_graficas/flujo_metodologico.png"):
    n = len(ETAPAS)
    alto = 1.4 + n * 2.12 + 0.6
    fig, ax = plt.subplots(figsize=(7, alto / 2))
    ax.set_xlim(0, 14)
    ax.set_ylim(0, alto)
    ax.axis("off")
    ax.text(7, alto - 0.6, "Flujo metodológico — Generalización cross-dataset", ha="center", va="center",
            fontsize=13, fontweight="bold")
    y = alto - 1.4
    for k, (tit, l1, l2) in enumerate(ETAPAS, start=1):
        ax.add_patch(FancyBboxPatch((1.2, y - 1.52), 11.6, 1.52, boxstyle="round,pad=0,rounding_size=0.12",
                                    linewidth=1.6, edgecolor="#1F4E79", facecolor="#EAF0F8"))
        ax.text(1.6, y - 0.42, f"{k}. {tit}", va="center", fontsize=11, fontweight="bold")
        ax.text(1.6, y - 0.86, l1, va="center", fontsize=9)
        ax.text(1.6, y - 1.2, l2, va="center", fontsize=9)
        if k < n:
            ax.annotate("", xy=(7, y - 2.08), xytext=(7, y - 1.6),
                        arrowprops=dict(arrowstyle="-|>", color="#404040", lw=1.6))
        y -= 2.12
    ax.text(7, 0.4, "Resultado del flujo: robustez de ViT-Base/16 y DenseNet-121 ante el cambio de dominio, "
            "con evidencia trazable.", ha="center", va="center", fontsize=8.5)
    fig.savefig(ruta, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return ruta


# --- formato (igual que Documentacion_Tecnica_10_Modelos.docx) ----------------
def nuevo_documento(plantilla):
    doc = Document(plantilla)
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    for rid, rel in list(doc.part.rels.items()):  # no arrastrar imagenes de la plantilla
        if rel.reltype.endswith("/image"):
            del doc.part.rels[rid]
    return doc


def run(p, texto, bold=None, italic=None, size=None, color=None, fuente=None):
    r = p.add_run(texto)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if fuente:
        r.font.name = fuente
    return r


def titulo(doc, h, n):
    p = doc.add_heading("", level=n)
    run(p, h, color=AZUL)
    return p


def parrafo(doc, texto, prefijo=None, italic=None, size=None, centrado=False, color=None):
    p = doc.add_paragraph()
    if prefijo:
        run(p, prefijo, bold=True)
    run(p, texto, italic=italic, size=size, color=color)
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def negrita(doc, texto, size=None):
    p = doc.add_paragraph()
    run(p, texto, bold=True, size=size)
    p.paragraph_format.keep_with_next = True
    return p


def formula(doc, *lineas):
    for linea in lineas:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent, pf.space_before, pf.space_after = Pt(17), Pt(4), Pt(8)
        run(p, linea, size=10, color=FORMULA, fuente="Consolas")


def archivo(doc, rutas, codigo):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent, pf.space_before, pf.space_after = Pt(11.35), Pt(2), Pt(8)
    run(p, "Archivo: ", bold=True, size=8.5, color=GRIS)
    run(p, rutas, size=8.5, color=GRIS, fuente="Consolas")
    run(p, "\n")
    run(p, "Código: ", bold=True, size=8.5, color=GRIS)
    run(p, codigo, size=8.5, color=GRIS, fuente="Consolas")


def leyenda(doc, texto):
    parrafo(doc, texto, italic=True, size=9, centrado=True, color=RGBColor(0x66, 0x66, 0x66))


def imagenes(doc, rutas, ancho):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    for r in rutas:
        p.add_run().add_picture(r, width=Cm(ancho))


def _hijo(padre, tag, **attrs):
    el = padre.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        padre.append(el)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


def _ancho_texto(s, pt):
    return sum(1.3 if c.isupper() else 1 for c in s) * pt * 13  # dxa aprox. por caracter


def tabla(doc, encabezado, filas, resaltar=()):
    doc.paragraphs[-1].paragraph_format.keep_with_next = True  # el titulo no queda solo al pie
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezado))
    tblPr = t._tbl.tblPr
    _hijo(tblPr, "w:tblW", w=5000, type="pct")
    _hijo(tblPr, "w:jc", val="center")
    _hijo(tblPr, "w:tblLayout", type="fixed")
    mar = _hijo(tblPr, "w:tblCellMar")
    for lado, v in (("top", MAR_TB), ("bottom", MAR_TB), ("left", MAR_LR), ("right", MAR_LR)):
        _hijo(mar, f"w:{lado}", w=v, type="dxa")
    pt = 9.5
    for ri, valores in enumerate([encabezado] + [list(f) for f in filas]):
        tr = t.rows[ri]._tr
        _hijo(tr.get_or_add_trPr(), "w:cantSplit")
        if ri == 0:
            _hijo(tr.get_or_add_trPr(), "w:tblHeader")
        for celda, valor in zip(t.rows[ri].cells, valores):
            tcPr = celda._tc.get_or_add_tcPr()
            if ri == 0 or ri - 1 in resaltar:
                _hijo(tcPr, "w:shd", val="clear", color="auto", fill=CABECERA if ri == 0 else RESALTE)
            _hijo(tcPr, "w:vAlign", val="center")
            p = celda.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before = pf.space_after = 0
            pf.line_spacing = 1.0
            run(p, str(valor), bold=ri == 0 or ri - 1 in resaltar, size=pt,
                color=RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else None)
    # anchos segun contenido: el encabezado no se corta a mitad de palabra
    sec = doc.sections[0]
    total = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)
    filas_txt = [encabezado] + [[str(v) for v in f] for f in filas]
    n = len(encabezado)
    need = [max(max(_ancho_texto(w, pt) for w in (fila[c].split() or [""])) for fila in filas_txt) for c in range(n)]
    want = [max(_ancho_texto(fila[c][:45], pt) if i else need[c] for i, fila in enumerate(filas_txt))
            for c in range(n)]
    base = [x + 2 * MAR_LR for x in need]
    if sum(base) > total:  # no cabe: se reduce la letra de toda la tabla
        escala = (total - 2 * MAR_LR * n) / sum(need)
        base = [x * escala + 2 * MAR_LR for x in need]
        for r in t._tbl.xpath(".//w:r"):
            r.get_or_add_rPr().get_or_add_sz().val = Pt(max(7, pt * escala))
    extra = total - sum(base)
    anchos = [int(b + extra * w / sum(want)) for b, w in zip(base, want)]
    anchos[-1] += total - sum(anchos)
    for gc, w in zip(t._tbl.tblGrid.gridCol_lst, anchos):
        gc.w = Twips(w)
    for fila in t.rows:
        for celda, w in zip(fila.cells, anchos):
            _hijo(celda._tc.get_or_add_tcPr(), "w:tcW", w=w, type="dxa")
    return t


def f2(x):
    return f"{x:.2f}"


# --- documento -----------------------------------------------------------------
GLOSARIO = [
    ("Accuracy (exactitud)", "Porcentaje de imágenes clasificadas correctamente sobre el total: (TP+TN)/total × 100."),
    ("AdamW", "Optimizador que ajusta los pesos con la regla de Adam y aplica el weight decay por separado del "
              "gradiente."),
    ("AUC", "Área bajo la curva ROC. Mide qué tan bien el modelo ordena las imágenes por probabilidad, sin "
            "depender de un umbral."),
    ("Augmentación x2", "Variante en la que cada imagen de entrenamiento aparece dos veces por epoch con "
                        "transformaciones distintas. Este documento usa la corrida sin x2."),
    ("Backbone", "Parte de la red que extrae las características de la imagen; sobre ella va la cabeza de "
                 "clasificación."),
    ("Backbone congelado", "Backbone con pesos fijos (sin gradiente) y en modo eval(); no cambia durante el "
                           "entrenamiento en el dominio objetivo."),
    ("Batch size", "Cantidad de imágenes que el modelo procesa juntas antes de actualizar los pesos (8 en "
                   "ViT-Base/16, 16 en DenseNet-121)."),
    ("BatchNorm", "Capa que normaliza las activaciones con estadísticas de los datos; con el backbone en modo "
                  "eval() esas estadísticas no se actualizan."),
    ("Cabeza de clasificación", "Última capa lineal del modelo, que convierte las características en las dos "
                                "clases; es la única parte que se entrena aquí."),
    ("Casi-duplicado", "Foto prácticamente idéntica a otra (por ejemplo, tomada en ráfaga). Se detectan con "
                       "detectar_duplicados.py y se mantienen en el mismo fold."),
    ("Checkpoint", "Archivo (.pt) con los pesos de un modelo. Aquí cada fold parte del checkpoint de Mendeley "
                   "del mismo índice de fold."),
    ("Clase positiva", "Clase que se toma como referencia para contar TP, FP, FN y TN; en este documento es "
                       "fresco, salvo cuando se indica lo contrario."),
    ("CNN", "Red neuronal convolucional: arquitectura que procesa la imagen con filtros locales (DenseNet-121)."),
    ("Cross-dataset", "Evaluación de un modelo entrenado en un conjunto de datos sobre otro conjunto distinto."),
    ("CrossEntropyLoss", "Función de pérdida de clasificación que penaliza la probabilidad baja asignada a la "
                         "clase correcta."),
    ("Curva ROC", "Gráfica de la tasa de verdaderos positivos frente a la de falsos positivos para todos los "
                  "umbrales posibles."),
    ("Data augmentation", "Transformaciones aleatorias aplicadas solo a las imágenes de entrenamiento (recorte, "
                          "volteo, color, blur, rotación, erasing)."),
    ("Desviación estándar muestral", "Medida de dispersión de los 5 F1 de fold, calculada con n − 1 en el "
                                     "denominador. Un valor bajo indica un modelo estable."),
    ("Dominio fuente", "Conjunto con el que se entrenaron los modelos base: Mendeley (8,000 imágenes)."),
    ("Dominio objetivo", "Conjunto nuevo al que se transfiere el modelo: Trucha o Rohu."),
    ("Early stopping", "Regla que detiene el entrenamiento cuando val_F1 no mejora durante un número de epochs "
                       "(paciencia = 8)."),
    ("Epoch", "Una pasada completa del modelo por todas las imágenes de entrenamiento."),
    ("Especificidad", "De las imágenes que no son de la clase evaluada, porcentaje que el modelo descartó "
                      "correctamente: TN/(TN+FP) × 100."),
    ("F1 macro", "Promedio simple del F1 de las dos clases (fresco y no_fresco). Es la métrica principal de este "
                 "documento."),
    ("F1-Score", "Media armónica de Precisión y Recall de una clase: 2 × P × R / (P + R)."),
    ("FN (falso negativo)", "Imagen de la clase evaluada que el modelo asignó a la otra clase."),
    ("Fold", "Cada una de las 5 particiones de un dominio. En la corrida k, el fold k se evalúa y los otros 4 "
             "entrenan la cabeza."),
    ("FP (falso positivo)", "Imagen de la otra clase que el modelo asignó a la clase evaluada."),
    ("Fuga de información", "Situación en la que datos de evaluación influyen en el entrenamiento, lo que infla "
                            "las métricas; los grupos de casi-duplicados la evitan."),
    ("Gap (brecha de error)", "Aumento del error al pasar del dominio fuente al objetivo: Error_objetivo − "
                              "Error_fuente (ver 2.6)."),
    ("ImageNet", "Conjunto de imágenes de referencia con el que se preentrenan los modelos; también define la "
                 "normalización de las imágenes."),
    ("Inferencia", "Uso del modelo ya entrenado para predecir imágenes, sin modificar sus pesos."),
    ("Learning rate", "Tamaño del paso con que el optimizador actualiza los pesos (aquí 1e-4)."),
    ("Linear probe (congelado)", "Método de transferencia que congela el backbone y entrena solo la cabeza de "
                                 "clasificación en el dominio objetivo."),
    ("Matriz de confusión", "Tabla 2 × 2 que cruza la clase real (filas) con la clase predicha (columnas); de "
                            "ella salen TP, FP, FN y TN."),
    ("Matriz pooled", "Suma celda a celda de las 5 matrices de fold de un modelo; suma 1,150 imágenes en Trucha "
                      "y 1,261 en Rohu."),
    ("MPS", "Metal Performance Shaders: aceleración por GPU de PyTorch en equipos Apple."),
    ("OOF (out-of-fold)", "Predicción que recibe una imagen de la cabeza entrenada con los otros 4 folds, que "
                          "no la vio en entrenamiento."),
    ("Paciencia", "Número de epochs seguidos sin mejora de val_F1 que se toleran antes del early stopping."),
    ("Parámetros", "Cantidad de pesos del modelo, en millones (M); los entrenables aquí son solo los de la "
                   "cabeza."),
    ("Pooled", "Resultado calculado una sola vez sobre todas las predicciones OOF concatenadas, no como promedio "
               "de folds."),
    ("Precisión", "De las imágenes que el modelo asignó a una clase, porcentaje correcto: TP/(TP+FP) × 100."),
    ("PyTorch", "Biblioteca de deep learning usada para entrenar y evaluar los modelos."),
    ("RBE (reducción de brecha de error)", "Cuánto reduce un modelo el Gap frente al baseline DenseNet-121, en "
                                           "porcentaje (ver 2.6)."),
    ("Recall (sensibilidad)", "De las imágenes de una clase, porcentaje que el modelo detectó: TP/(TP+FN) × 100."),
    ("Scheduler coseno", "Regla que reduce el learning rate siguiendo una curva coseno a lo largo de los epochs "
                         "(CosineAnnealingLR)."),
    ("scikit-learn", "Biblioteca de machine learning usada para la partición (StratifiedKFold, "
                     "StratifiedGroupKFold) y las métricas."),
    ("Semilla (random_state)", "Número fijo (42) que hace reproducibles los procesos aleatorios, como la "
                               "partición de los datos."),
    ("Softmax", "Función que convierte las salidas del modelo en probabilidades por clase que suman 1."),
    ("StratifiedGroupKFold", "Partición en k folds que mantiene la proporción de clases y deja cada grupo de "
                             "casi-duplicados completo en un solo fold (Rohu)."),
    ("StratifiedKFold", "Partición en k folds que mantiene en cada fold la misma proporción de clases (Trucha)."),
    ("timm", "Biblioteca de PyTorch con las arquitecturas y pesos preentrenados de los modelos."),
    ("TN (verdadero negativo)", "Imagen de la otra clase que el modelo también asignó a la otra clase."),
    ("TP (verdadero positivo)", "Imagen de la clase evaluada que el modelo asignó a esa clase."),
    ("Transformer", "Arquitectura basada en mecanismos de atención; en visión divide la imagen en parches (ViT)."),
    ("TRD (tasa de retención del desempeño)", "Porcentaje del F1 de Mendeley que el modelo conserva en el dominio "
                                              "objetivo (ver 2.6)."),
    ("Umbral", "Probabilidad de corte para decidir la clase; la matriz de confusión usa 0.5."),
    ("val_F1", "F1 macro sobre el fold de evaluación, medido al final de cada epoch para elegir el mejor."),
    ("Validación cruzada", "Esquema en el que cada fold se evalúa una vez, de modo que todas las imágenes se "
                           "evalúan fuera de muestra."),
    ("ViT", "Vision Transformer: aplica un Transformer a parches de 16 × 16 píxeles de la imagen; aquí se usa "
            "ViT-Base/16."),
    ("Weight decay", "Penalización que reduce los pesos grandes para evitar sobreajuste (aquí 1e-4)."),
    ("Zero-shot", "Evaluación sin ningún entrenamiento en el dominio objetivo; no es el método de este documento."),
]


def generar(datos, plantilla):
    doc = nuevo_documento(plantilla)
    p = parrafo(doc, "", centrado=True)
    run(p, "Documentación Técnica Detallada", bold=True, size=26, color=AZUL)
    p = parrafo(doc, "", centrado=True)
    run(p, "Clasificación de Frescura de Pescado — Generalización Cross-Dataset (Trucha, Rohu)", size=14,
        color=RGBColor(0x44, 0x44, 0x44))
    doc.add_paragraph()
    parrafo(doc, "Congelado (linear probe) desde el checkpoint de Mendeley — ViT-Base/16 y DenseNet-121, "
                 "5-fold, sin augmentación x2", size=11, centrado=True)
    parrafo(doc, "Cálculos por fold y pooled — con fórmulas, gráficas y datos trazables a los archivos fuente",
            italic=True, size=11, centrado=True)
    doc.add_paragraph()

    # 1. Metodologia ---------------------------------------------------------------
    titulo(doc, "1. Metodología — flujo en 7 etapas", 1)
    parrafo(doc, "El experimento mide cuánto conservan ViT-Base/16 y DenseNet-121, entrenados sobre Mendeley "
                 "(dominio fuente; ver Documentacion_Tecnica_10_Modelos.docx), al pasar a dos dominios objetivo: "
                 "Trucha y Rohu.")
    imagenes(doc, [figura_flujo()], 12)
    leyenda(doc, "Flujo metodológico — generalización cross-dataset de ViT-Base/16 y DenseNet-121 (7 etapas, de "
                 "la preparación de los dominios objetivo a los indicadores de robustez).")
    parrafo(doc, "Cada etapa del flujo se documenta en una subsección: la etapa k corresponde a la sección 1.k. "
                 "El glosario de fórmulas está en la sección 2, el detalle por modelo y dominio en la sección 3 "
                 "y el glosario de términos técnicos en el Anexo B.")

    titulo(doc, "1.1 Etapa 1 — Dominios objetivo y preparación de datos", 2)
    parrafo(doc, "Trucha: 1,150 imágenes (OneDrive), 481 fresco / 669 no_fresco, dominio objetivo original del "
                 "proyecto. Se excluyen 46 imágenes marcadas como dudosas y se incluyen 82 rescatadas tras "
                 "revisión manual.", prefijo="Trucha. ")
    parrafo(doc, "Rohu: 1,261 imágenes, 422 fresco / 839 no_fresco, dominio objetivo evaluado bajo el mismo "
                 "protocolo y el mismo pipeline de entrenamiento que Trucha. Se excluyen 22 imágenes dudosas y "
                 "109 falsos positivos confirmados del recorte de ojo (escama, aleta o borrosa), y se incluyen "
                 "48 rescatadas manualmente.", prefijo="Rohu. ")
    parrafo(doc, "Las clases no están balanceadas en ninguno de los dos dominios (Mendeley sí lo está, 50/50); "
                 "por eso la métrica central es el F1 macro, que pesa igual a fresco y no_fresco (ver 2.3).")

    titulo(doc, "1.2 Etapa 2 — Validación cruzada 5-fold en cada dominio", 2)
    parrafo(doc, "Cada dominio se divide en 5 folds con semilla 42. En cada corrida, 4 folds se usan para "
                 "entrenar la cabeza y el fold restante para evaluar; cada imagen se evalúa una sola vez.")
    formula(doc, "Fold i:  Evaluación_i (≈20%)  |  Entrenamiento_i (≈80%)   [i = 1..5]")
    parrafo(doc, "StratifiedKFold(n_splits=5, shuffle=True, random_state=42): 920 imágenes de entrenamiento y "
                 "230 de evaluación por fold, con la misma proporción de clases en cada fold.", prefijo="Trucha. ")
    parrafo(doc, "StratifiedGroupKFold(n_splits=5, shuffle=True, random_state=42), agrupando las fotos casi "
                 "idénticas (tomadas en ráfaga, detectadas con detectar_duplicados.py) para que una foto y su casi "
                 "gemela nunca queden una en entrenamiento y otra en evaluación. Por eso el tamaño del fold de "
                 "evaluación varía entre 228 y 284 imágenes. Se detectaron 108 grupos de 2 o más fotos casi "
                 "idénticas (el más grande con 15); cada una de las demás imágenes es su propio grupo.",
            prefijo="Rohu. ")
    parrafo(doc, "a diferencia de Mendeley (68/12/20), aquí no hay un fold de validación separado del de "
                 "prueba: el mismo fold sirve para elegir el mejor epoch (1.4) y para reportar sus métricas "
                 "(1.5). Esto puede dar un resultado algo optimista, por lo que las cifras se leen como "
                 "comparación entre los dos modelos bajo el mismo protocolo, no como desempeño esperado en "
                 "producción.", prefijo="Limitación: ")

    titulo(doc, "1.3 Etapa 3 — Configuración de los 2 modelos (congelado / linear probe)", 2)
    parrafo(doc, "Cada fold parte del checkpoint de Mendeley (8,000 imágenes, 10 modelos evaluados en "
                 "Documentacion_Tecnica_10_Modelos.docx) del mismo índice de fold (el fold 3 de Trucha parte "
                 "del checkpoint del fold 3, etc.). El backbone se congela (pesos fijos, sin gradiente, en modo "
                 "eval() para que BatchNorm no actualice sus estadísticas) y solo se entrena la cabeza de "
                 "clasificación sobre el dominio objetivo. Es entrenamiento real (hay backpropagation y "
                 "actualización de pesos en la cabeza), no una evaluación zero-shot.")
    tabla(doc, ["Modelo", "Familia", "Parámetros totales", "Parámetros entrenables (cabeza)",
                "Tamaño de imagen", "Batch size"],
          [(n, fam, par, f"{cab:,}", f"{img}×{img}", b) for _, n, fam, par, img, b, cab in MODELOS])
    parrafo(doc, "optimizador AdamW sobre la cabeza (learning rate 1e-4, weight decay 1e-4), "
                 "CosineAnnealingLR, CrossEntropyLoss, hasta 30 epochs con early stopping (paciencia 8 epochs sin "
                 "mejora de val_F1).", prefijo="Hiperparámetros (finetune_target_x2.py, finetune_eval_fold): ")
    parrafo(doc, "cada imagen de entrenamiento aparece una sola vez por epoch, con la augmentación "
                 "estocástica estándar de get_transforms_v3 (recorte, flip, color, blur, rotación, erasing). La "
                 "evaluación nunca se aumenta.", prefijo="Sin augmentación x2: ")
    negrita(doc, "Arquitectura del sistema y entorno de ejecución")
    parrafo(doc, "Flujo de datos entre las cuatro etapas del pipeline (ingesta, entrenamiento, evaluación y "
                 "reportes) y los scripts que las ejecutan — desde las imágenes de Trucha y Rohu y los "
                 "checkpoints de Mendeley hasta los resultados documentados.")
    imagenes(doc, ["4_graficas/arquitectura_sistema.png"], 16)
    negrita(doc, "Especificaciones técnicas del entorno de ejecución:")
    tabla(doc, ["Componente", "Detalle"], [
        ("Hardware", "PC con procesador Intel Core i9 y 64 GB de RAM"),
        ("Deep learning", "PyTorch + torchvision (entrenamiento, tensores, transforms)"),
        ("Model zoo", "timm — ViT-Base/16 y DenseNet-121"),
        ("ML utilitario", "scikit-learn — StratifiedKFold, StratifiedGroupKFold, confusion_matrix, roc_curve, "
                          "roc_auc_score"),
        ("Datos y cálculo", "pandas, numpy — CSV, matrices, métricas"),
        ("Visualización", "matplotlib — matrices de confusión y curvas ROC"),
        ("Documentación", "python-docx — generación de este documento (generar_documentacion_tecnica.py)"),
    ])

    titulo(doc, "1.4 Etapa 4 — Entrenamiento de la cabeza y selección del mejor epoch", 2)
    parrafo(doc, "Por cada combinación (modelo, dominio, fold) — 20 en total — se entrena la cabeza hasta 30 "
                 "epochs. En cada epoch se mide val_F1 sobre el fold de evaluación; se guardan las predicciones "
                 "del epoch con mejor val_F1, que son las que se reportan (no las del último epoch).")
    negrita(doc, "Corridas experimentales por modelo y dominio")
    filas = []
    for clave, nombre, *_ in MODELOS:
        for dom, dominio in DOMINIOS:
            d = datos[dom, clave]
            filas.append((nombre, dominio, "5", sum(f["epochs"] for f in d["folds"]), "5", f"{d['pooled'].sum():,}"))
    tabla(doc, ["Modelo", "Dominio", "Corridas (folds)", "Epochs ejecutados", "Evaluaciones", "Predicciones OOF"],
          filas)
    parrafo(doc, "Epochs ejecutados = suma de los 5 folds (máximo 30 por fold; menos de 150 indica early "
                 "stopping). Corresponden a la corrida que guardó matrices de confusión y curvas ROC "
                 "(linear_probe_confmat_roc.py). Esa corrida no registró val_F1 por epoch; el detalle epoch por "
                 "epoch de la sección 3.x.1 viene de una corrida de verificación con 10 folds (mismo código, "
                 "datos y semilla), que no reemplaza las cifras presentadas.", prefijo="Nota. ")

    titulo(doc, "1.5 Etapa 5 — Evaluación por fold", 2)
    parrafo(doc, "Con las predicciones del mejor epoch de cada fold se arma la matriz de confusión del fold y se "
                 "aplican las fórmulas del glosario (sección 2). La probabilidad softmax de la clase no_fresco "
                 "se usa para la curva ROC y el AUC del fold.")
    parrafo(doc, "Métricas por fold con su matriz de confusión y fórmulas instanciadas: sección 3.x de cada "
                 "modelo y dominio («Resultados por fold»); curvas ROC por fold: «Curvas ROC por fold».")

    titulo(doc, "1.6 Etapa 6 — Consolidación OOF / pooled", 2)
    parrafo(doc, "Las predicciones de evaluación de los 5 folds, llamadas predicciones out-of-fold (OOF, ver "
                 "glosario 2.5), se concatenan: cada imagen aparece una sola vez, en el fold donde fue evaluada. "
                 "Sobre ese conjunto se calcula UNA sola vez la matriz de confusión y las métricas:")
    formula(doc, "Matriz_pooled = Matriz_fold1 + Matriz_fold2 + Matriz_fold3 + Matriz_fold4 + Matriz_fold5")
    negrita(doc, "Cálculo de la matriz pooled y de los números generales de cada modelo")
    parrafo(doc, "las predicciones del mejor epoch de cada fold forman una matriz 2×2 (fresco / no_fresco).",
            prefijo="Paso 1 — Matriz de cada fold: ")
    parrafo(doc, "se suman celda a celda las 5 matrices de fold (confusion_matrix_pooled.csv).",
            prefijo="Paso 2 — Matriz pooled: ")
    parrafo(doc, "para no_fresco se invierten los roles (su TP es el TN de fresco).",
            prefijo="Paso 3 — Métricas por clase desde la matriz pooled: ")
    formula(doc, "Precisión_fresco = TP / (TP+FP)     Recall_fresco = TP / (TP+FN)",
            "Precisión_no_fresco = TN / (TN+FN)     Recall_no_fresco = TN / (TN+FP)",
            "F1_clase = 2 × Precisión × Recall / (Precisión + Recall)")
    parrafo(doc, "", prefijo="Paso 4 — Números generales (macro = promedio simple de las dos clases):")
    formula(doc, "Accuracy = (TP+TN) / N × 100",
            "Precisión macro = (Precisión_fresco + Precisión_no_fresco) / 2;  igual para Recall y F1")
    parrafo(doc, "el AUC pooled se calcula con las probabilidades OOF concatenadas (roc_curve_pooled.csv), no "
                 "con la matriz. La desviación estándar del F1 es la desviación muestral (n − 1) de los 5 F1 de "
                 "fold.", prefijo="Paso 5 — Números que no salen de la matriz: ")
    clave, nombre = MODELOS[0][:2]
    m = metricas(datos["trucha", clave]["pooled"])
    negrita(doc, f"Ejemplo instanciado — {nombre}, Trucha")
    formula(doc,
            f"Matriz pooled: TP={m['tp']}  FN={m['fn']}  FP={m['fp']}  TN={m['tn']}   "
            f"(TP+FN = {m['tp'] + m['fn']}; FP+TN = {m['fp'] + m['tn']})",
            f"Precisión_fresco = {m['tp']}/({m['tp']}+{m['fp']}) = {f2(m['p_f'])}     Recall_fresco = "
            f"{m['tp']}/{m['tp'] + m['fn']} = {f2(m['r_f'])}     F1_fresco = {f2(m['f1_f'])}",
            f"Precisión_no_fresco = {m['tn']}/({m['tn']}+{m['fn']}) = {f2(m['p_n'])}     Recall_no_fresco = "
            f"{m['tn']}/{m['fp'] + m['tn']} = {f2(m['r_n'])}     F1_no_fresco = {f2(m['f1_n'])}",
            f"Accuracy = ({m['tp']}+{m['tn']})/{m['n']} × 100 = {f2(m['acc'])}",
            f"Precisión macro = ({f2(m['p_f'])}+{f2(m['p_n'])})/2 = {f2(m['prec'])}     F1 macro = "
            f"({f2(m['f1_f'])}+{f2(m['f1_n'])})/2 = {f2(m['f1'])}")
    negrita(doc, "Matriz pooled y números generales por modelo y dominio")
    filas = []
    for dom, dominio in DOMINIOS:
        for clave, nombre, *_ in MODELOS:
            m = metricas(datos[dom, clave]["pooled"])
            filas.append((nombre, dominio, m["tp"], m["fn"], m["fp"], m["tn"], f2(m["acc"]), f2(m["prec"]),
                          f2(m["f1"])))
    tabla(doc, ["Modelo", "Dominio", "TP", "FN", "FP", "TN", "Accuracy", "Precisión macro", "F1 macro"], filas,
          resaltar=(0, 2))
    parrafo(doc, "TP/FN = imágenes fresco predichas fresco/no_fresco; FP/TN = imágenes no_fresco predichas "
                 "fresco/no_fresco. El detalle con la suma de las 5 matrices de fold está en cada sección 3.x.",
            prefijo="Nota. ")

    titulo(doc, "1.7 Etapa 7 — Comparación e indicadores de robustez", 2)
    filas = []
    for dom, dominio in DOMINIOS:
        for clave, nombre, *_ in MODELOS:
            d = datos[dom, clave]
            m = metricas(d["pooled"])
            sd = statistics.stdev(metricas(f["cm"])["f1"] for f in d["folds"])
            filas.append((nombre, dominio, f2(m["acc"]), f2(m["prec"]), f2(m["f1"]), f2(sd), f"{d['auc']:.4f}",
                          f2(d["f1_oficial"])))
    tabla(doc, ["Modelo", "Dominio", "Accuracy", "Precisión", "F1 macro (pooled)", "Desv.std F1 (5 folds)",
                "AUC (pooled)", "F1 oficial (media 5 folds)"], filas, resaltar=(0, 2))
    parrafo(doc, "los valores pooled provienen de la corrida de matrices de confusión y ROC "
                 "(linear_probe_confmat_roc.py, misma configuración, reentrenada), por eso difieren levemente "
                 "del F1 oficial (finetune_target_x2.py y experimento_rohu.py, 3_resultados/"
                 "{trucha,rohu}/congelado_sinx2_resumen.csv). Los indicadores de robustez usan el F1 oficial.",
            prefijo="Nota: ", italic=True, size=9.5)
    negrita(doc, "F1 oficial por fold (corrida de 3_resultados/, base de los indicadores de robustez)")
    filas = []
    for dom, dominio in DOMINIOS:
        df = pd.read_csv(f"3_resultados/{dom}/congelado_sinx2_por_fold.csv")
        for clave, nombre, *_ in MODELOS:
            for r in df[df.model == clave].itertuples():
                n_eval = r.n_val
                filas.append((nombre, dominio, f"Fold {r.fold}", r.n_train, n_eval, f2(r.f1), f2(r.accuracy),
                              r.epochs_run, f"{r.tiempo_s:.1f}"))
            filas.append((nombre, dominio, "Media", "", "", f2(datos[dom, clave]["f1_oficial"]), "", "", ""))
    tabla(doc, ["Modelo", "Dominio", "Fold", "N entrenamiento", "N evaluación", "F1", "Accuracy", "Epochs",
                "Tiempo (s)"], filas, resaltar=(5, 11, 17, 23))
    parrafo(doc, "corrida de finetune_target_x2.py (Trucha) y experimento_rohu.py (Rohu): mismo "
                 "método y semilla que la corrida de matrices de confusión, pero ejecutada aparte; por eso el F1 "
                 "de cada fold difiere levemente del de las secciones 3.x. Logs: " + LOG["trucha"] + " y "
                 + LOG["rohu"] + ".", prefijo="Nota: ", italic=True, size=9.5)
    negrita(doc, "Indicadores de robustez (fórmulas en 2.6)")
    filas = []
    for dom, dominio in DOMINIOS:
        for clave, nombre, *_ in MODELOS:
            trd, gap, rbe = indicadores(datos, dom, clave)
            filas.append((nombre + (" (baseline)" if clave == "densenet121" else ""), dominio,
                          f2(datos[dom, clave]["f1_oficial"]), f2(F1_FUENTE[clave]), f2(trd), f"{gap:.4f}",
                          f2(rbe)))
    tabla(doc, ["Modelo", "Dominio", "F1 objetivo", "F1 fuente (Mendeley)", "TRD (%)", "Gap",
                "RBE vs. DenseNet-121 (%)"], filas, resaltar=(0, 2))
    negrita(doc, "Lectura")
    parrafo(doc, f"ViT-Base/16 rinde por encima de DenseNet-121 en los dos dominios, consistente con su "
                 f"desempeño en Mendeley ({F1_FUENTE['vit_base_patch16_224']} vs. {F1_FUENTE['densenet121']}).")
    parrafo(doc, "Rohu rinde por encima de Trucha en los 2 modelos.")
    parrafo(doc, "ViT-Base/16 retiene más desempeño que DenseNet-121 en ambos dominios (TRD más alto, Gap más "
                 "chico) y su RBE es positivo: transfiere mejor que el baseline CNN.")

    # 2. Glosario ------------------------------------------------------------------
    titulo(doc, "2. Glosario de fórmulas", 1)
    parrafo(doc, "Todas las métricas parten de una matriz de confusión 2x2 (fresco vs no_fresco). TP/TN/FP/FN se "
                 "definen tomando cada clase como positiva por turno (one-vs-rest), y las versiones reportadas "
                 "son el promedio macro entre las 2 clases.")
    titulo(doc, "2.1 Qué significa cada sigla (TP, FP, FN, TN)", 2)
    parrafo(doc, "Estas 4 siglas son la base de TODAS las fórmulas de este documento. Se definen tomando una "
                 "clase como \"positiva\" (la que se está evaluando en esa fila de la tabla) y la otra como "
                 "\"negativa\":")
    tabla(doc, ["Sigla", "Nombre completo", "Qué significa"], [
        ("TP", "Verdadero Positivo (True Positive)", "Era de esta clase Y el modelo acertó (la predijo)"),
        ("FP", "Falso Positivo (False Positive)",
         "NO era de esta clase, pero el modelo la predijo como si lo fuera (falsa alarma)"),
        ("FN", "Falso Negativo (False Negative)", "SÍ era de esta clase, pero el modelo predijo la otra (no la "
                                                  "detectó)"),
        ("TN", "Verdadero Negativo (True Negative)",
         "NO era de esta clase Y el modelo tampoco la predijo (acertó al descartarla)"),
    ])
    parrafo(doc, "Ejemplo con la clase \"fresco\": TP = imágenes fresco que el modelo dijo fresco. FP = imágenes "
                 "no_fresco que el modelo dijo (por error) fresco. FN = imágenes fresco que el modelo dijo (por "
                 "error) no_fresco. TN = imágenes no_fresco que el modelo dijo correctamente no_fresco.")
    parrafo(doc, "Importante: en la fila \"no_fresco\" de cualquier tabla de este documento, los roles se "
                 "invierten — el TP de no_fresco es el TN de fresco, y así sucesivamente.", italic=True, size=9.5)
    titulo(doc, "2.2 Fórmulas de las métricas", 2)
    tabla(doc, ["Métrica", "Fórmula", "Interpretación"], [
        ("Accuracy", "(TP+TN) / (TP+TN+FP+FN) × 100", "Global — no depende de qué clase es 'positiva'"),
        ("Precisión", "TP / (TP+FP) × 100", "De lo que predije como esta clase, ¿cuánto acerté?"),
        ("Recall (Sensibilidad)", "TP / (TP+FN) × 100", "De lo que era esta clase, ¿cuánto detecté?"),
        ("F1-Score", "2×Precisión×Recall / (Precisión+Recall)",
         "Media armónica — penaliza si Precisión o Recall es bajo"),
        ("Especificidad", "TN / (TN+FP) × 100", "De lo que NO era esta clase, ¿cuánto rechacé bien?"),
        ("F1 (estándar/binario)", "F1 de UNA sola clase (la 'positiva')", "No usado como métrica final"),
        ("F1-macro (usado aquí)", "(F1_fresco + F1_no_fresco) / 2",
         "Promedio simple de las 2 clases — trata a ambas por igual"),
        ("AUC", "Área bajo la curva ROC (0 a 1)", "Evalúa TODOS los umbrales de decisión, no solo 50%"),
    ])
    titulo(doc, "2.3 Diferencia F1 vs F1-macro", 2)
    parrafo(doc, "El F1 'estándar' evalúa una sola clase. El F1-macro, usado en todo este proyecto, promedia el F1 "
                 "de fresco y de no_fresco por igual. En Trucha y Rohu la diferencia importa más que en Mendeley "
                 "porque las clases no están balanceadas (no_fresco es mayoría): el F1 de fresco es varios puntos "
                 "menor que el de no_fresco en los 4 casos (ver 3.x, resultado pooled).")
    negrita(doc, "Aclaración importante — la columna \"F1\" de las tablas de este documento:")
    parrafo(doc, "En TODAS las tablas de este documento, la columna llamada simplemente \"F1\" es siempre el "
                 "F1-MACRO — nunca el F1 de una sola clase.")
    titulo(doc, "2.4 Por qué el AUC puede diferir del Accuracy en el mismo modelo", 2)
    parrafo(doc, "El Accuracy usa un solo corte de decisión (softmax > 50% → una clase u otra). El AUC evalúa la "
                 "calidad de la probabilidad en TODOS los cortes posibles — por eso un modelo puede tener Accuracy "
                 "moderado en el corte de 50% pero AUC alto si sus probabilidades igual ordenan bien los casos.")
    titulo(doc, "2.5 Predicciones out-of-fold (OOF) y resultado pooled", 2)
    parrafo(doc, "OOF (out-of-fold, \"fuera del fold\") es la predicción que recibe una imagen cuando forma parte "
                 "del fold de evaluación. Esa predicción la hace la cabeza entrenada con los otros 4 folds, que "
                 "nunca vio esa imagen en entrenamiento.")
    parrafo(doc, "Como cada imagen está en la evaluación de exactamente un fold, al terminar los 5 folds se tiene "
                 "exactamente una predicción OOF por imagen: 1,150 en Trucha y 1,261 en Rohu.")
    parrafo(doc, "Resultado pooled: son las métricas calculadas UNA sola vez sobre todas las predicciones OOF "
                 "concatenadas. No es el promedio de las métricas de los 5 folds: el F1 y el AUC no son lineales y "
                 "se recalculan sobre el total. La matriz de confusión pooled sí equivale a la suma celda a celda "
                 "de las 5 matrices de fold.")
    parrafo(doc, "Media ± desviación estándar por fold: es el complemento del resultado pooled. Resume cuánto varía "
                 "el F1 cuando cambia la partición; una desviación baja indica un modelo estable.")
    titulo(doc, "2.6 Indicadores de robustez ante cambio de dominio (TRD, Gap, RBE)", 2)
    parrafo(doc, "Indicadores del plan de tesis (Dimensión 2 — Robustez ante el Cambio de Dominio):")
    tabla(doc, ["Indicador", "Fórmula", "Interpretación"], [
        ("TRD — Tasa de Retención del Desempeño", "TRD = (F1_objetivo / F1_fuente) × 100",
         "Porcentaje del desempeño de Mendeley que el modelo conserva en el dominio objetivo. 100% = sin caída"),
        ("Gap — Brecha de Error", "Error = 1 − F1/100;  Gap = Error_objetivo − Error_fuente",
         "Cuánto aumenta el error al pasar del dominio fuente al objetivo"),
        ("RBE — Reducción de Brecha de Error", "RBE = (Gap_baseline − Gap_modelo) / Gap_baseline × 100",
         "Cuánto reduce el modelo la brecha frente a DenseNet-121 (baseline CNN). RBE > 0 = transfiere mejor"),
    ])
    parrafo(doc, "Para el propio baseline, Gap_baseline = Gap_modelo, por lo que RBE = 0.0 exacto — es una "
                 "comparación contra sí mismo, no una medición independiente; se muestra igual para no dejar la "
                 "celda vacía.", italic=True, size=9.5)

    # 3. Detalle por modelo ------------------------------------------------------
    titulo(doc, "3. Detalle por modelo", 1)
    for i, (clave, nombre, fam, par, *_) in enumerate(MODELOS, start=1):
        titulo(doc, f"3.{i}. {nombre}", 2)
        resumen = "   ".join(f"F1 macro {dominio} = {f2(metricas(datos[dom, clave]['pooled'])['f1'])}"
                               for dom, dominio in DOMINIOS)
        parrafo(doc, f"F1 macro Mendeley (fuente) = {F1_FUENTE[clave]}   {resumen}   Parámetros = {par}",
                prefijo="")
        titulo(doc, f"3.{i}.1 Entrenamiento — val_F1 por epoch", 3)
        archivo(doc, f"3_resultados/{{trucha,rohu}}/confmat_roc/resumen_confmat_roc.csv (corrida presentada) y "
                     f"3_resultados/{{trucha,rohu}}/verificacion_10fold/{clave}/epochs_fold{{1..10}}.csv "
                     f"(corrida de verificación)",
                "finetune_target_x2.py — finetune_eval_fold(): mide val_F1 al final de cada epoch y corta cuando "
                "pasan 8 epochs seguidos sin mejora.")
        parrafo(doc, "La corrida presentada (5 folds) solo guardó, por fold, el total de epochs ejecutados y "
                     "las métricas del mejor epoch. Con eso se arma el resumen por fold. El detalle epoch por "
                     "epoch viene de la corrida de verificación (10 folds): mismo código, datos, semilla e "
                     "hiperparámetros, pero con 10 particiones, por lo que sus folds no son los mismos que los "
                     "de la corrida presentada y sus cifras no la reemplazan.", prefijo="Nota. ")
        for dom, dominio in DOMINIOS:
            d = datos[dom, clave]
            negrita(doc, f"{dominio} — corrida presentada (5 folds): resumen por fold")
            tabla(doc, ["Fold", "Epoch elegido (mejor val_F1)", "val_F1 en ese epoch", "Total epochs ejecutados"],
                  [(f"Fold {k}", epoch_elegido(f["epochs"]), f2(metricas(f["cm"])["f1"]), f["epochs"])
                   for k, f in enumerate(d["folds"], start=1)])
            parrafo(doc, "el entrenamiento se detiene cuando pasan 8 epochs seguidos sin mejora, así que si "
                         "el fold paró antes de 30 epochs el mejor fue el epoch total − 8. Si llegó a 30, el mejor "
                         "está entre el 22 y el 30 y no quedó registrado. val_F1 en ese epoch es el F1 del fold, "
                         "porque el mismo fold elige el epoch y se evalúa (ver 1.2).", prefijo="Epoch elegido: ",
                    size=9.5)
            png, folds = curvas_epoch(dom, clave, nombre, dominio)
            r = pd.read_csv(f"{VERIF.format(dom=dom)}/resumen.csv")
            r = r[(r.model == clave) & (r.fold != "pooled")]
            negrita(doc, f"{dominio} — corrida de verificación (10 folds): val_F1 por epoch")
            imagenes(doc, [png], 15)
            leyenda(doc, f"{nombre} — {dominio} — val_F1 por epoch en la corrida de verificación. El punto marca "
                         f"el epoch elegido (mejor val_F1 de ese fold).")
            filas = []
            for k, (e, (_, fila)) in enumerate(zip(folds, r.iterrows()), start=1):
                mejor = e.loc[e.val_f1.idxmax()]
                assert int(mejor.epoch) == int(float(fila.epoch_elegido)) and len(e) == int(float(fila.epochs_run))
                filas.append((f"Fold {k}", int(mejor.epoch), f2(mejor.val_f1), len(e)))
            tabla(doc, ["Fold", "Epoch elegido", "val_F1 en ese epoch", "Total epochs ejecutados"], filas)
            parrafo(doc, f"Media del F1 de los 10 folds = {f2(r.test_f1.mean())} ± {f2(r.test_f1.std())}. "
                         f"Cada fold k parte del checkpoint de Mendeley del fold ((k − 1) mod 5) + 1.", size=9.5)
            negrita(doc, "Detalle completo — todos los epochs de cada fold (fila resaltada = epoch elegido):")
            for k, e in enumerate(folds, start=1):
                parrafo(doc, f"Fold {k}", italic=True)
                tabla(doc, ["Epoch", "val_F1", "Mejor hasta ahora", "Paciencia"],
                      [(int(x.epoch), f2(x.val_f1), f2(x.mejor_hasta_ahora), int(x.paciencia))
                       for x in e.itertuples()],
                      resaltar=(int(e.val_f1.idxmax()),))
        sub = 1
        for dom, dominio in DOMINIOS:
            d = datos[dom, clave]
            ruta = f"3_resultados/{dom}/confmat_roc/{clave}"
            cm_png, roc_png = graficas_folds(datos, dom, clave, nombre, dominio)

            sub += 1
            titulo(doc, f"3.{i}.{sub} {dominio} — resultados por fold", 3)
            archivo(doc, f"{ruta}/confusion_matrix_fold{{1..5}}.csv  y  "
                         f"3_resultados/{dom}/confmat_roc/resumen_confmat_roc.csv (filas de este modelo)",
                    "linear_probe_confmat_roc.py — process_model(): por cada fold llama a finetune_eval_fold() "
                    "(finetune_target_x2.py) y guarda la matriz y la curva ROC del mejor epoch.")
            parrafo(doc, "Las predicciones del mejor epoch de cada fold forman la matriz de confusión del fold; "
                         "se aplican las mismas fórmulas del glosario (sección 2):")
            formula(doc, "Precisión = TP / (TP+FP) × 100     Recall = TP / (TP+FN) × 100",
                    "F1 = 2 × Precisión × Recall / (Precisión + Recall)",
                    "F1 macro del fold = (F1_fresco + F1_no_fresco) / 2")
            negrita(doc, "Qué fold se evalúa y cuáles entrenan en cada fila — para no confundirlos:")
            filas = []
            for k, f in enumerate(d["folds"], start=1):
                m = metricas(f["cm"])
                filas.append((f"Fold {k}", f"Fold {k}", ",".join(str(j) for j in range(1, 6) if j != k), m["n"],
                              f2(m["acc"]), f2(m["prec"]), f2(m["rec"]), f2(m["f1"]), f"{f['auc']:.4f}",
                              f["epochs"]))
            tabla(doc, ["Corrida", "Fold de EVALUACIÓN", "Folds de ENTRENAMIENTO", "N", "Accuracy", "Precisión",
                        "Recall", "F1", "AUC", "Epochs"], filas)
            m = metricas(d["folds"][0]["cm"])
            negrita(doc, f"Ejemplo instanciado — Fold 1 (matriz: TP={m['tp']}, FP={m['fp']}, FN={m['fn']}, "
                         f"TN={m['tn']}):")
            formula(doc, f"F1_fresco = 2×{f2(m['p_f'])}×{f2(m['r_f'])}/({f2(m['p_f'])}+{f2(m['r_f'])}) = "
                         f"{f2(m['f1_f'])}",
                    f"F1_no_fresco = 2×{f2(m['p_n'])}×{f2(m['r_n'])}/({f2(m['p_n'])}+{f2(m['r_n'])}) = "
                    f"{f2(m['f1_n'])}",
                    f"F1 macro Fold 1 = ({f2(m['f1_f'])}+{f2(m['f1_n'])})/2 = {f2(m['f1'])}   "
                    f"(valor en la tabla: {f2(m['f1'])})")
            imagenes(doc, [cm_png], 15)
            leyenda(doc, f"{nombre} — {dominio} — matriz de confusión de cada uno de los 5 folds (mejor epoch "
                         f"evaluado sobre el fold de evaluación).")

            sub += 1
            titulo(doc, f"3.{i}.{sub} {dominio} — curvas ROC por fold", 3)
            archivo(doc, f"{ruta}/roc_curve_fold{{1..5}}.csv  (columnas fpr, tpr, threshold)",
                    "linear_probe_confmat_roc.py — roc_curve() y roc_auc_score() sobre la probabilidad softmax "
                    "del mejor epoch.")
            imagenes(doc, [roc_png], 10)
            leyenda(doc, f"{nombre} — {dominio} — curva ROC de cada fold individual (color) y curva pooled "
                         f"(negro).")

            sub += 1
            titulo(doc, f"3.{i}.{sub} {dominio} — resultado pooled, cómo se calculó cada número", 3)
            archivo(doc, f"{ruta}/confusion_matrix_pooled.csv, roc_curve_pooled.csv",
                    "linear_probe_confmat_roc.py — final de process_model(): concatena las predicciones de los "
                    "5 folds y aplica compute_metrics() (de metrics.py) una sola vez.")
            negrita(doc, "Matriz de confusión pooled = suma de las 5 matrices de fold:")
            cms = [f["cm"] for f in d["folds"]]
            mp = metricas(d["pooled"])
            for etiqueta, (a, b), clave_m in (("TP (fresco→fresco)", (0, 0), "tp"),
                                              ("FN (fresco→no_fresco)", (0, 1), "fn"),
                                              ("FP (no_fresco→fresco)", (1, 0), "fp"),
                                              ("TN (no_fresco→no_fresco)", (1, 1), "tn")):
                formula(doc, f"{etiqueta}: {'+'.join(str(c[a, b]) for c in cms)} = {mp[clave_m]}")
            tp, fn, fp, tn = mp["tp"], mp["fn"], mp["fp"], mp["tn"]
            parrafo(doc, "El F1 macro es el promedio del F1 de cada clase, y el F1 de una clase se obtiene de su "
                         "Precisión y su Recall, que a su vez dependen de sus conteos TP, FP y FN. La matriz de "
                         "confusión es la tabla que contiene exactamente esos conteos para las dos clases. Con "
                         f"fresco como clase positiva se leen TP = {tp}, FN = {fn}, FP = {fp} y TN = {tn}; para "
                         f"no_fresco los roles se invierten (su TP es {tn}, su FP es {fn} y su FN es {fp}). Así se "
                         "calculan las métricas de cada clase de la tabla siguiente y luego su promedio macro, en "
                         "el que fresco y no_fresco pesan lo mismo.",
                    prefijo="¿Por qué se calcula desde la matriz de confusión? ")
            parrafo(doc, f"La matriz también muestra en qué dirección se equivoca {nombre}, algo que el Accuracy "
                         f"por sí solo no indica: {fn} imágenes fresco se clasificaron como no_fresco y {fp} "
                         f"imágenes no_fresco como fresco. Se usa la matriz pooled porque reúne las "
                         f"{mp['n']:,} predicciones OOF, cada imagen una sola vez; por eso el F1 macro reportado "
                         f"se calcula desde ella y no promediando el F1 de los 5 folds. El AUC es la única métrica "
                         f"que no sale de la matriz, porque usa las probabilidades y no un único umbral (ver 2.4).")
            tabla(doc, ["Clase", "TP", "FP", "FN", "TN", "Precisión", "Recall", "F1", "Especificidad"], [
                ("fresco", tp, fp, fn, tn, f2(mp["p_f"]), f2(mp["r_f"]), f2(mp["f1_f"]), f2(mp["r_n"])),
                ("no_fresco", tn, fn, fp, tp, f2(mp["p_n"]), f2(mp["r_n"]), f2(mp["f1_n"]), f2(mp["r_f"])),
            ])
            formula(doc, f"F1 macro = ({f2(mp['f1_f'])} + {f2(mp['f1_n'])}) / 2 = {f2(mp['f1'])}",
                    f"Accuracy = ({tp}+{tn}) / {mp['n']} × 100 = {f2(mp['acc'])}")
            imagenes(doc, [f"4_graficas/{dom}_sinx2/{clave}/confusion_matrix_pooled.png",
                           f"4_graficas/{dom}_sinx2/{clave}/roc_curve_pooled.png"], 7.0)
            leyenda(doc, f"{nombre} — {dominio} — matriz de confusión pooled (izq.) y curva ROC pooled (der.), "
                         f"AUC = {d['auc']:.4f}.")

    # Anexo ------------------------------------------------------------------------
    titulo(doc, "Anexo A. Trazabilidad de archivos fuente", 1)
    parrafo(doc, "Todos los números de este documento provienen de archivos reales, sin datos simulados.")
    tabla(doc, ["Qué", "Dónde"], [
        ("Código de entrenamiento y evaluación",
         "1_codigo/finetune_target_x2.py, experimento_rohu.py, linear_probe_confmat_roc.py"),
        ("Generación de este documento y de las gráficas por fold", "1_codigo/generar_documentacion_tecnica.py"),
        ("Logs de entrenamiento", f"{LOG['trucha']}, {LOG['rohu']}, 2_logs/rohu_confmat_roc.log"),
        ("F1 oficial por fold y resumen", "3_resultados/{trucha,rohu}/congelado_sinx2_{por_fold,resumen}.csv"),
        ("Matrices de confusión y curvas ROC (por fold y pooled)",
         "3_resultados/{trucha,rohu}/confmat_roc/{modelo}/*.csv y resumen_confmat_roc.csv"),
        ("val_F1 por epoch (corrida de verificación, 10 folds)",
         "3_resultados/{trucha,rohu}/verificacion_10fold/{modelo}/epochs_fold{1..10}.csv y resumen.csv"),
        ("Gráficas", "4_graficas/{trucha,rohu}_sinx2/{modelo}/"),
        ("F1 macro de Mendeley (dominio fuente)",
         "Mendeley_ViT_DenseNet/3_resultados_test_por_modelo/metricas_pooled_por_modelo.csv"),
    ])
    titulo(doc, "Anexo B. Glosario de términos técnicos", 1)
    parrafo(doc, "Términos técnicos usados en este documento, en orden alfabético. Las fórmulas de las métricas "
                 "están en la sección 2.")
    tabla(doc, ["Término", "Definición"], sorted(GLOSARIO, key=lambda t: t[0].lower()))
    doc.save(OUT)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--plantilla", default=PLANTILLA)
    args = ap.parse_args()
    generar(cargar(), args.plantilla)
    print(f"Guardado en {OUT}")
